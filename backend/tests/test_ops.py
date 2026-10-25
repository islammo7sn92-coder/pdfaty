"""Operation-level tests: every test checks the produced file itself."""
import json
import os
import zipfile

import pikepdf
import pypdfium2 as pdfium
import pytest


def pages(path, pw=""):
    with pikepdf.open(path, password=pw) as p:
        return len(p.pages)


# ---------------------------------------------------------------- opening / info
def test_info(run, samples):
    ctx, _ = run("info", {}, [samples["english"]])
    assert ctx.result["pages"] == 6
    assert ctx.result["encrypted"] is False


def test_info_encrypted_requires_password(run, samples, user_error):
    with pytest.raises(user_error) as e:
        run("info", {}, [samples["encrypted"]])
    assert e.value.code == "password_required"
    with pytest.raises(user_error) as e:
        run("info", {"password": "nope"}, [samples["encrypted"]])
    assert e.value.code == "wrong_password"
    ctx, _ = run("info", {"password": "user123"}, [samples["encrypted"]])
    assert ctx.result["encrypted"] and ctx.result["pages"] == 6


def test_invalid_pdf(run, samples, user_error):
    with pytest.raises(user_error) as e:
        run("info", {}, [samples["bad"]])
    assert e.value.code == "invalid_pdf"


# ---------------------------------------------------------------- page management
def test_merge(run, samples, text_of):
    ctx, outs = run("merge", {}, [samples["english"], samples["arabic"]])
    assert pages(outs[0]) == 6 + pages(samples["arabic"])
    t = text_of(outs[0])
    assert "Hello world page 1" in t and "عقد توريد" in t


def test_merge_keeps_form_fields(run, samples):
    _, outs = run("merge", {}, [samples["form"], samples["english"]])
    with pikepdf.open(outs[0]) as p:
        names = {str(f.get("/T", "")) for f in p.Root.AcroForm.Fields}
    assert {"name", "agree", "city"} <= names


def test_split_ranges_and_each(run, samples, text_of):
    ctx, outs = run("split", {"mode": "ranges", "ranges": "1-2;3;4-6"}, [samples["english"]])
    pdfs = [o for o in outs if o.endswith(".pdf")]
    assert [pages(o) for o in pdfs] == [2, 1, 3]
    assert "page 3" in text_of(pdfs[1])
    assert any(o.endswith(".zip") for o in outs)
    ctx, outs = run("split", {"mode": "each"}, [samples["english"]])
    assert len([o for o in outs if o.endswith(".pdf")]) == 6


def test_extract_delete_rotate(run, samples, text_of):
    _, outs = run("extract_pages", {"pages": "2,4"}, [samples["english"]])
    assert pages(outs[0]) == 2 and "page 4" in text_of(outs[0])
    _, outs = run("delete_pages", {"pages": "1-3"}, [samples["english"]])
    assert pages(outs[0]) == 3 and "page 1" not in text_of(outs[0])
    _, outs = run("rotate", {"pages": "1", "angle": 90}, [samples["english"]])
    with pikepdf.open(outs[0]) as p:
        assert int(p.pages[0].obj.Rotate) == 90 and int(p.pages[1].obj.get("/Rotate", 0)) == 0


def test_organize_reorder_insert_blank_duplicate(run, samples, text_of):
    plan = [{"src": 0, "page": 2}, {"src": 1, "page": 0, "rotate": 180}, {"blank": True}, {"src": 0, "page": 0},
            {"src": 0, "page": 0}]
    _, outs = run("organize", {"pages": plan}, [samples["english"], samples["arabic"]])
    with pikepdf.open(outs[0]) as p:
        assert len(p.pages) == 5
        assert int(p.pages[1].obj.Rotate) == 180
    doc = pdfium.PdfDocument(outs[0])
    texts = []
    for i in range(len(doc)):
        tp = doc[i].get_textpage()
        texts.append(tp.get_text_range())
    assert "page 3" in texts[0] and "عقد" in texts[1] and texts[2].strip() == "" and "page 1" in texts[3] and "page 1" in texts[4]


def test_blank_detection(run, samples):
    ctx, _ = run("detect_blank", {}, [samples["english"]])
    assert ctx.result["blank_pages"] == [5]
    ctx, outs = run("remove_blank", {}, [samples["english"]])
    assert pages(outs[0]) == 5


def test_bad_range(run, samples, user_error):
    with pytest.raises(user_error):
        run("extract_pages", {"pages": "9-12"}, [samples["english"]])


# ---------------------------------------------------------------- watermark / background / numbers
def test_watermark_text_add_analyze_remove(run, samples, text_of):
    _, outs = run("watermark_add", {"text": "سري CONFIDENTIAL", "opacity": 0.3, "rotation": 45, "pages": "1-3",
                                    "color": "#ff0000", "size": 48}, [samples["english"]])
    wm = outs[0]
    ctx, _ = run("watermark_analyze", {}, [wm])
    ids = [c["id"] for c in ctx.result["candidates"] if c["source"] == "pdfaty"]
    assert ids and ctx.result["candidates"][0]["page_count"] == 3
    _, outs = run("watermark_remove", {"ids": ids}, [wm])
    ctx, _ = run("watermark_analyze", {}, [outs[0]])
    assert not [c for c in ctx.result["candidates"] if c["source"] == "pdfaty"]
    assert "Hello world page 1" in text_of(outs[0])


def test_watermark_image_tiled_and_preview(run, samples):
    ctx, outs = run("watermark_add", {"type": "image", "tile": True, "opacity": 0.2, "scale": 0.2},
                    [samples["english"], samples["sig"]])
    assert pages(outs[0]) == 6
    ctx, outs = run("watermark_add", {"text": "DRAFT", "preview_page": 1}, [samples["english"]])
    assert pages(outs[0]) == 1


def test_watermark_auto_remove_artifact(run, samples):
    _, outs = run("watermark_add", {"text": "DRAFT"}, [samples["english"]])
    _, outs2 = run("watermark_remove", {}, [outs[0]])
    ctx, _ = run("watermark_analyze", {}, [outs2[0]])
    assert not ctx.result["candidates"]


def test_watermark_remove_nothing(run, samples, user_error):
    with pytest.raises(user_error) as e:
        run("watermark_remove", {}, [samples["english"]])
    assert e.value.code == "nothing_to_remove"


def test_external_watermark_text_detected(run, samples, tmp_path):
    # a "foreign" watermark: rotated grey text drawn as normal content on every page
    from reportlab.pdfgen import canvas

    p = str(tmp_path / "foreign.pdf")
    c = canvas.Canvas(p)
    for i in range(3):
        c.drawString(72, 700, f"Body text {i}")
        c.saveState()
        c.setFillGray(0.8)
        c.translate(300, 400)
        c.rotate(45)
        c.setFont("Helvetica", 60)
        t = c.beginText(0, 0)
        t.setTextOrigin(0, 0)
        t.textOut("SAMPLE")
        c.drawText(t)
        c.restoreState()
        c.showPage()
    c.save()
    ctx, _ = run("watermark_analyze", {}, [p])
    text_c = [c for c in ctx.result["candidates"] if c["source"] == "text"]
    assert text_c and text_c[0]["page_count"] == 3
    _, outs = run("watermark_remove", {"ids": [text_c[0]["id"]]}, [p])
    from conftest import pdf_text

    t = pdf_text(outs[0])
    assert "SAMPLE" not in t and "Body text 1" in t


def test_background_add_remove(run, samples):
    _, outs = run("background_add", {"color": "#fff3c4"}, [samples["english"]])
    doc = pdfium.PdfDocument(outs[0])
    px = doc[0].render(scale=0.2).to_pil().convert("RGB").getpixel((3, 3))
    assert px[2] < 230  # yellowish
    ctx, _ = run("watermark_analyze", {}, [outs[0]])
    assert any(c["kind"] == "background" for c in ctx.result["candidates"])
    _, outs2 = run("background_remove", {}, [outs[0]])
    doc = pdfium.PdfDocument(outs2[0])
    px = doc[0].render(scale=0.2).to_pil().convert("RGB").getpixel((3, 3))
    assert px == (255, 255, 255)
    _, outs3 = run("background_add", {"type": "image", "fit": "cover", "opacity": 0.5}, [samples["english"], samples["photo"]])
    assert pages(outs3[0]) == 6


def test_header_footer_numbers(run, samples, text_of):
    _, outs = run("header_footer", {"numbering": True, "number_format": "page_n_of", "start": 1, "skip_first": True,
                                    "slots": {"tc": "{title} - {date}"}, "title": "Report"}, [samples["english"]])
    t = text_of(outs[0])
    assert "Page 1 of 5" in t and "Page 5 of 5" in t and "Report" in t
    _, outs = run("header_footer", {"numbering": True, "number_format": "ar_n_of", "digits": "arabic"}, [samples["english"]])
    assert pages(outs[0]) == 6


# ---------------------------------------------------------------- editing
def test_edit_overlay_text_image_shapes_link(run, samples, text_of):
    items = [
        {"type": "text", "page": 0, "x": 72, "y": 120, "w": 300, "h": 40, "text": "Added text مرحبا", "size": 16, "color": "#0000ff"},
        {"type": "image", "page": 0, "x": 300, "y": 600, "w": 150, "h": 60, "file": 1},
        {"type": "rect", "page": 0, "x": 50, "y": 50, "w": 100, "h": 50, "stroke": "#ff0000"},
        {"type": "arrow", "page": 0, "x1": 100, "y1": 300, "x2": 200, "y2": 350},
        {"type": "whiteout", "page": 1, "x": 60, "y": 60, "w": 200, "h": 30},
        {"type": "link", "page": 0, "x": 72, "y": 70, "w": 200, "h": 20, "url": "https://example.com"},
        {"type": "image", "pages": "2-3", "x": 400, "y": 700, "w": 120, "h": 50, "file": 1},
    ]
    _, outs = run("edit_overlay", {"items": items}, [samples["english"], samples["sig"]])
    assert "Added text" in text_of(outs[0])
    with pikepdf.open(outs[0]) as p:
        links = [a for a in p.pages[0].obj.Annots if a.Subtype == "/Link"]
        assert links and str(links[0].A.URI) == "https://example.com"


def test_edit_link_rejects_js(run, samples, user_error):
    with pytest.raises(user_error):
        run("edit_overlay", {"items": [{"type": "link", "page": 0, "x": 1, "y": 1, "w": 5, "h": 5, "url": "javascript:alert(1)"}]},
            [samples["english"]])


def test_replace_existing_text(run, samples, text_of):
    ctx, _ = run("text_objects", {"page": 0}, [samples["english"]])
    objs = ctx.result["objects"]
    target = [o for o in objs if "INV-1000" in o["text"]][0]
    _, outs = run("replace_text", {"edits": [{"page": 0, "ids": [target["index"]], "text": "Invoice number INV-2000 total 99.00 USD",
                                              "x": target["rect"][0], "y": target["rect"][1], "w": 400, "h": 20,
                                              "size": target["size"], "baseline": target["baseline"]}]}, [samples["english"]])
    t = text_of(outs[0])
    assert "INV-2000" in t and "INV-1000" not in t.split("\n")[0:3].__str__()
    first_page = pdfium.PdfDocument(outs[0])[0].get_textpage().get_text_range()
    assert "INV-1000" not in first_page and "INV-2000" in first_page


def test_replace_arabic_text(run, samples):
    ctx, _ = run("text_objects", {"page": 0}, [samples["arabic"]])
    objs = ctx.result["objects"]
    assert objs
    t0 = [o for o in objs if "25000" in o["text"]]
    assert t0, [o["text"] for o in objs][:30]
    o = t0[0]
    _, outs = run("replace_text", {"edits": [{"page": 0, "ids": [o["index"]], "text": "26000", "x": o["rect"][0], "y": o["rect"][1],
                                              "w": o["rect"][2] + 20, "h": o["rect"][3], "size": o["size"], "baseline": o["baseline"]}]},
                  [samples["arabic"]])
    first_page = pdfium.PdfDocument(outs[0])[0].get_textpage().get_text_range()
    assert "26000" in first_page and "25000" not in first_page


# ---------------------------------------------------------------- annotations
def test_annotations_roundtrip(run, samples):
    items = [
        {"type": "highlight", "page": 0, "rects": [[72, 75, 150, 16]], "contents": "check"},
        {"type": "underline", "page": 0, "rects": [[72, 95, 150, 16]]},
        {"type": "strikeout", "page": 0, "rects": [[72, 95, 150, 16]]},
        {"type": "note", "page": 0, "x": 400, "y": 100, "contents": "ملاحظة مهمة"},
        {"type": "freetext", "page": 0, "x": 300, "y": 200, "w": 200, "h": 60, "text": "تعليق نصي"},
        {"type": "ink", "page": 0, "paths": [[[100, 500], [120, 520], [140, 505]]]},
        {"type": "arrow", "page": 0, "x1": 100, "y1": 600, "x2": 200, "y2": 650},
        {"type": "square", "page": 0, "x": 50, "y": 700, "w": 80, "h": 40},
        {"type": "circle", "page": 0, "x": 150, "y": 700, "w": 80, "h": 40},
        {"type": "stamp", "page": 0, "x": 350, "y": 400, "w": 160, "h": 50, "text": "معتمد", "stamp": "approved"},
    ]
    _, outs = run("annotate", {"items": items, "author": "Reviewer"}, [samples["english"]])
    ctx, _ = run("annotations_list", {}, [outs[0]])
    anns = ctx.result["annotations"]
    types = {a["type"] for a in anns}
    assert {"Highlight", "Underline", "StrikeOut", "Text", "FreeText", "Ink", "Line", "Square", "Circle", "Stamp"} <= types
    note = [a for a in anns if a["type"] == "Text"][0]
    assert note["contents"] == "ملاحظة مهمة" and note["author"] == "Reviewer"
    hl = [a for a in anns if a["type"] == "Highlight"][0]
    _, outs2 = run("annot_edit", {"actions": [{"id": note["id"], "action": "update", "contents": "تم التعديل"},
                                              {"id": hl["id"], "action": "delete"}]}, [outs[0]])
    ctx, _ = run("annotations_list", {}, [outs2[0]])
    anns2 = ctx.result["annotations"]
    assert not [a for a in anns2 if a["type"] == "Highlight"]
    assert [a for a in anns2 if a["type"] == "Text"][0]["contents"] == "تم التعديل"


# ---------------------------------------------------------------- forms
def test_forms_list_fill_flatten(run, samples, text_of):
    ctx, _ = run("forms_list", {}, [samples["form"]])
    fields = {f["name"]: f for f in ctx.result["fields"]}
    assert fields["name"]["type"] == "text" and fields["agree"]["type"] == "checkbox"
    assert fields["gender"]["type"] == "radio" and fields["city"]["type"] == "combo"
    vals = {"name": "محمد أحمد", "agree": True, "gender": "female", "city": "Giza"}
    _, outs = run("forms_fill", {"values": vals}, [samples["form"]])
    ctx, _ = run("forms_list", {}, [outs[0]])
    got = {f["name"]: f["value"] for f in ctx.result["fields"]}
    assert got["name"] == "محمد أحمد" and got["gender"] == "female" and got["city"] == "Giza" and got["agree"] not in ("", "Off")
    _, outs = run("forms_fill", {"values": vals, "flatten": True}, [samples["form"]])
    with pikepdf.open(outs[0]) as p:
        assert "/AcroForm" not in p.Root
    assert "Giza" in text_of(outs[0])


def test_forms_create(run, samples):
    specs = [
        {"type": "text", "name": "full name", "page": 0, "x": 72, "y": 200, "w": 200, "h": 22},
        {"type": "checkbox", "name": "ok", "page": 0, "x": 72, "y": 240, "w": 16, "h": 16, "value": True},
        {"type": "radio", "name": "choice", "export": "a", "page": 0, "x": 72, "y": 270, "w": 14, "h": 14},
        {"type": "radio", "name": "choice", "export": "b", "page": 0, "x": 100, "y": 270, "w": 14, "h": 14, "value": True},
        {"type": "dropdown", "name": "city", "page": 0, "x": 72, "y": 300, "w": 120, "h": 20, "options": ["Cairo", "Giza"]},
        {"type": "signature", "name": "sig", "page": 0, "x": 72, "y": 340, "w": 150, "h": 40},
        {"type": "date", "name": "date", "page": 0, "x": 250, "y": 200, "w": 100, "h": 22},
    ]
    _, outs = run("forms_create", {"fields": specs}, [samples["english"]])
    ctx, _ = run("forms_list", {}, [outs[0]])
    f = {x["name"]: x for x in ctx.result["fields"]}
    assert f["full_name"]["type"] == "text" and f["ok"]["type"] == "checkbox" and f["choice"]["type"] == "radio"
    assert f["choice"]["value"] == "b" and f["city"]["type"] == "combo" and f["sig"]["type"] == "signature"
    _, outs2 = run("forms_fill", {"values": {"full_name": "Test User", "city": "Giza"}}, [outs[0]])
    ctx, _ = run("forms_list", {}, [outs2[0]])
    assert {x["name"]: x["value"] for x in ctx.result["fields"]}["full_name"] == "Test User"


# ---------------------------------------------------------------- protection / redaction
def test_protect_unprotect(run, samples, text_of, user_error):
    _, outs = run("protect", {"user_password": "s3cret", "owner_password": "own3r", "permissions": {"print": False, "copy": False}},
                  [samples["english"]])
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(outs[0])
    with pikepdf.open(outs[0], password="s3cret") as p:
        assert p.is_encrypted and not p.allow.print_lowres and not p.allow.extract
        assert p.encryption.R == 6
    with pytest.raises(user_error) as e:
        run("unprotect", {"password": "s3cret"}, [outs[0]])
    assert e.value.code == "owner_password_required"
    _, outs2 = run("unprotect", {"password": "own3r"}, [outs[0]])
    with pikepdf.open(outs2[0]) as p:
        assert not p.is_encrypted
    assert "Hello world" in text_of(outs2[0])


def test_restricted_pdf_is_honoured(run, samples, user_error):
    with pytest.raises(user_error) as e:
        run("watermark_add", {"text": "X"}, [samples["restricted"]])
    assert e.value.code == "permission_denied"
    with pytest.raises(user_error):
        run("pdf_to_text", {}, [samples["restricted"]])
    _, outs = run("watermark_add", {"text": "X", "password": "owner456"}, [samples["restricted"]])
    assert outs


def test_redaction_search_and_apply(run, samples, text_of):
    ctx, _ = run("redact_search", {"terms": ["sales@example.com"], "presets": ["national_id"]}, [samples["arabic"]])
    hits = ctx.result["matches"]
    assert len(hits) >= 2
    _, outs = run("redact_apply", {"terms": ["sales@example.com"], "presets": ["national_id"]}, [samples["arabic"]])
    t = text_of(outs[0])
    assert "sales@example.com" not in t and "29801011234567" not in t
    assert "عقد توريد" in t  # unrelated text kept
    with pikepdf.open(outs[0]) as p:
        raw = b"".join(p.pages[0].obj.Contents.read_bytes() if not isinstance(p.pages[0].obj.Contents, pikepdf.Array)
                       else [c.read_bytes() for c in p.pages[0].obj.Contents])
    assert b"sales@example" not in raw


def test_redaction_area_on_scanned_page(run, samples, text_of):
    # scanned page: the image must be rewritten so pixels under the box are gone
    _, outs = run("redact_apply", {"areas": [{"page": 0, "rect": [50, 50, 300, 80]}]}, [samples["scanned"]])
    doc = pdfium.PdfDocument(outs[0])
    img = doc[0].render(scale=1).to_pil().convert("L")
    region = img.crop((100, 70, 300, 120))
    assert max(region.getdata()) < 40  # black
    with pikepdf.open(outs[0]) as p:
        # the original full-page image must no longer be referenced unmodified
        imgs = [o for o in p.objects if isinstance(o, pikepdf.Stream) and o.get("/Subtype") == "/Image"]
        for im in imgs:
            pil = pikepdf.PdfImage(im).as_pil_image().convert("L")
            w, h = pil.size
            if w > 1000:
                # big image => must contain the black box burned in
                sx = w / 595.28
                assert max(pil.crop((int(110 * sx), int(80 * sx), int(280 * sx), int(110 * sx))).getdata()) < 60


def test_redaction_rasterize_mode(run, samples, text_of):
    ctx, outs = run("redact_apply", {"terms": ["INV-1000"], "mode": "rasterize"}, [samples["english"]])
    assert "INV-1000" not in text_of(outs[0])
    assert ctx.result["rasterized_pages"] == [1]
    assert ctx.result["verified"]


def test_sanitize(run, samples):
    ctx, outs = run("sanitize", {}, [samples["arabic"]])
    with pikepdf.open(outs[0]) as p:
        assert "/Metadata" not in p.Root and not dict(p.docinfo)


# ---------------------------------------------------------------- compression
def test_compress_levels(run, tmp_path):
    from PIL import Image
    import numpy as np
    from reportlab.pdfgen import canvas

    arr = (np.random.RandomState(1).rand(1600, 1600, 3) * 255).astype("uint8")
    im = Image.fromarray(arr).filter(__import__("PIL.ImageFilter", fromlist=["x"]).GaussianBlur(3))
    ip = str(tmp_path / "big.png")
    im.save(ip)
    p = str(tmp_path / "heavy.pdf")
    c = canvas.Canvas(p)
    c.drawString(72, 800, "Heavy document")
    c.drawImage(ip, 72, 300, 300, 300)
    c.showPage()
    c.save()
    for level in ("light", "medium", "strong"):
        ctx, outs = run("compress", {"level": level}, [p])
        r = ctx.result
        assert r["after"] < r["before"] and r["valid"]
        assert os.path.getsize(outs[0]) == r["after"]
    from conftest import pdf_text

    assert "Heavy document" in pdf_text(outs[0])


# ---------------------------------------------------------------- conversion
def test_pdf_to_images_and_text(run, samples):
    ctx, outs = run("pdf_to_images", {"format": "jpg", "dpi": 72, "pages": "1-2"}, [samples["english"]])
    jpgs = [o for o in outs if o.endswith(".jpg")]
    assert len(jpgs) == 2 and zipfile.is_zipfile([o for o in outs if o.endswith(".zip")][0])
    ctx, outs = run("pdf_to_text", {}, [samples["arabic"]])
    t = open(outs[0], encoding="utf-8").read()
    assert "عقد توريد أجهزة" in t and "1,250,000" in t


def test_pdf_to_docx(run, samples):
    from docx import Document

    _, outs = run("pdf_to_docx", {}, [samples["arabic"]])
    d = Document(outs[0])
    body = "\n".join(p.text for p in d.paragraphs)
    cells = " ".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "عقد توريد أجهزة" in body
    assert "حاسب محمول" in cells or "حاسب محمول" in body


def test_pdf_to_xlsx_csv(run, samples):
    from openpyxl import load_workbook

    ctx, outs = run("pdf_to_xlsx", {}, [samples["english"]])
    wb = load_workbook(outs[0])
    values = [c.value for ws in wb for row in ws.iter_rows() for c in row]
    assert "Pen" in values and "Book" in values
    ctx, outs = run("pdf_to_csv", {"pages": "1"}, [samples["english"]])
    csvs = [o for o in outs if o.endswith(".csv")]
    assert csvs and "Pen" in open(csvs[0], encoding="utf-8-sig").read()


def test_pdf_to_pptx(run, samples):
    from pptx import Presentation

    _, outs = run("pdf_to_pptx", {"pages": "1-2"}, [samples["english"]])
    assert len(Presentation(outs[0]).slides) == 2
    _, outs = run("pdf_to_pptx", {"pages": "1", "mode": "editable"}, [samples["arabic"]])
    texts = [sh.text_frame.text for s in Presentation(outs[0]).slides for sh in s.shapes if sh.has_text_frame]
    assert any("عقد" in t for t in texts)


def test_extract_images(run, samples):
    ctx, outs = run("extract_images", {}, [samples["english"]])
    assert ctx.result["images"] >= 1


def test_images_to_pdf(run, samples):
    ctx, outs = run("images_to_pdf", {"size": "a4"}, [samples["photo"], samples["sig"]])
    assert pages(outs[0]) == 2
    with pikepdf.open(outs[0]) as p:
        assert abs(float(p.pages[0].mediabox[2]) - 841.89) < 1  # landscape A4 for a wide photo


def test_office_to_pdf(run, samples, text_of):
    ctx, outs = run("convert_to_pdf", {"combine": True}, [samples["docx"], samples["xlsx"], samples["txt"], samples["photo"]])
    t = text_of(outs[0])
    assert "خطاب رسمي" in t and "أجهزة" in t and "ملاحظات الاجتماع" in t
    assert pages(outs[0]) >= 4


def test_blank_pdf(run):
    _, outs = run("blank_pdf", {"size": "a4", "count": 3, "orientation": "landscape"}, [])
    with pikepdf.open(outs[0]) as p:
        assert len(p.pages) == 3 and float(p.pages[0].mediabox[2]) > float(p.pages[0].mediabox[3])


# ---------------------------------------------------------------- OCR
@pytest.mark.timeout(300)
def test_ocr_arabic_scanned(run, samples, text_of):
    assert not text_of(samples["scanned"]).strip()
    ctx, outs = run("ocr", {"languages": ["ara", "eng"], "text_output": True}, [samples["scanned"]])
    t = text_of(outs[0])
    assert "عقد" in t and "توريد" in t
    assert "29801011234567" in t.replace(" ", "") or "2980101" in t
    assert ctx.result["stats"][0]["confidence"] > 50
    txt = open([o for o in outs if o.endswith(".txt")][0], encoding="utf-8").read()
    assert "sales@example.com" in txt
    # original image kept: page still renders the same scan
    with pikepdf.open(outs[0]) as p:
        assert len(p.pages) == 1


@pytest.mark.timeout(300)
def test_ocr_skip_text_pages(run, samples):
    ctx, outs = run("ocr", {"languages": ["eng"]}, [samples["english"]])
    assert 1 in ctx.result["skipped_pages"]


@pytest.mark.timeout(300)
def test_ocr_from_image(run, samples, text_of):
    ctx, outs = run("ocr", {"languages": ["ara"]}, [samples["scan_png"]])
    assert "توريد" in text_of(outs[0])


# ---------------------------------------------------------------- compare
def test_compare_contract(run, samples):
    ctx, outs = run("compare", {"mode": "contract", "lang": "ar"}, [samples["arabic"], samples["arabic_v2"]])
    r = ctx.result
    assert r["total"] >= 3
    blob = json.dumps(r["changes"], ensure_ascii=False)
    for old, new in (("30", "45"), ("1,250,000", "1,350,000"), ("25000", "27000")):
        assert old in blob and new in blob
    assert r["important"] >= 3 and r["changes"][0]["important"]
    assert any("amount" in c["categories"] for c in r["changes"])
    assert outs[0].endswith(".pdf") and pages(outs[0]) >= 2


def test_compare_identical(run, samples):
    ctx, _ = run("compare", {}, [samples["english"], samples["english"]])
    assert ctx.result["total"] == 0 and ctx.result["similarity"] == 100


# ---------------------------------------------------------------- inspection / metadata
def test_inspect(run, samples):
    ctx, _ = run("inspect", {}, [samples["arabic"]])
    r = ctx.result
    assert r["fonts"] and all(f["embedded"] for f in r["fonts"])
    ids = {c["id"]: c["status"] for c in r["checks"]["pdfa"]}
    assert ids["fonts_embedded"] == "pass" and r["disclaimer"] == "indicative_only"
    ctx, _ = run("inspect", {}, [samples["english"]])
    assert any(not f["embedded"] for f in ctx.result["fonts"])  # Helvetica is not embedded


def test_metadata_bookmarks_attachments(run, samples):
    _, outs = run("metadata_set", {"metadata": {"title": "عنوان", "author": "PDFATY", "lang": "ar", "display_title": True}},
                  [samples["english"]])
    with pikepdf.open(outs[0]) as p:
        assert str(p.docinfo.Title) == "عنوان" and str(p.Root.Lang) == "ar"
    bms = [{"title": "مقدمة", "page": 0, "children": [{"title": "Sub", "page": 1}]}, {"title": "End", "page": 5}]
    _, outs2 = run("bookmarks_set", {"bookmarks": bms}, [outs[0]])
    ctx, _ = run("bookmarks_get", {}, [outs2[0]])
    assert ctx.result["bookmarks"][0]["title"] == "مقدمة" and ctx.result["bookmarks"][0]["children"][0]["page"] == 1
    _, outs3 = run("attachments_add", {}, [outs2[0], samples["txt"]])
    ctx, outs4 = run("attachments_extract", {}, [outs3[0]])
    assert open(outs4[0], encoding="utf-8").read().startswith("ملاحظات")
    _, outs5 = run("attachments_remove", {}, [outs3[0]])
    with pikepdf.open(outs5[0]) as p:
        assert not p.attachments


def test_flatten_annotations(run, samples):
    _, outs = run("annotate", {"items": [{"type": "freetext", "page": 0, "x": 100, "y": 100, "w": 200, "h": 40, "text": "FLATME"}]},
                  [samples["english"]])
    _, outs2 = run("flatten", {}, [outs[0]])
    from conftest import pdf_text

    with pikepdf.open(outs2[0]) as p:
        assert not [a for a in p.pages[0].obj.get("/Annots", []) if a.get("/Subtype") == "/FreeText"]
    assert "FLATME" in pdf_text(outs2[0])


# ---------------------------------------------------------------- digital signatures
def test_digital_signature_and_verify(run, samples, tmp_path):
    from app.ops.signing import make_test_certificate

    p12 = str(tmp_path / "me.p12")
    make_test_certificate(p12, "certpw", "Ahmed Test")
    ctx, outs = run("sign_certificate", {"cert_password": "certpw", "reason": "Approval", "location": "Cairo",
                                         "page": 0, "rect": [350, 700, 200, 60]}, [samples["english"], p12])
    assert ctx.result["field"] == "Signature1"
    ctx2, _ = run("signatures_verify", {}, [outs[0]])
    sigs = ctx2.result["signatures"]
    assert len(sigs) == 1 and sigs[0]["intact"] and sigs[0]["valid"]
    assert sigs[0]["trusted"] is False  # self-signed: never claimed as trusted
    assert "Test" in sigs[0]["signer"]


def test_digital_signature_wrong_password(run, samples, tmp_path, user_error):
    from app.ops.signing import make_test_certificate

    p12 = str(tmp_path / "me.p12")
    make_test_certificate(p12, "right")
    with pytest.raises(user_error) as e:
        run("sign_certificate", {"cert_password": "wrong"}, [samples["english"], p12])
    assert e.value.code == "bad_certificate"


def test_sign_existing_field(run, samples, tmp_path):
    from app.ops.signing import make_test_certificate

    p12 = str(tmp_path / "me.p12")
    make_test_certificate(p12, "pw")
    _, outs = run("forms_create", {"fields": [{"type": "signature", "name": "ClientSig", "page": 0, "x": 72, "y": 600, "w": 160, "h": 50}]},
                  [samples["english"]])
    ctx, outs2 = run("sign_certificate", {"cert_password": "pw", "field_name": "ClientSig"}, [outs[0], p12])
    assert ctx.result["field"] == "ClientSig"
    ctx2, _ = run("signatures_verify", {}, [outs2[0]])
    assert ctx2.result["signatures"][0]["field"] == "ClientSig" and ctx2.result["signatures"][0]["valid"]


def test_digital_signature_arabic_name_always_valid(run, samples, tmp_path):
    from app.ops.signing import make_test_certificate

    p12 = str(tmp_path / "ar.p12")
    make_test_certificate(p12, "pw", "أحمد علي")
    ctx, outs = run("sign_certificate", {"cert_password": "pw", "page": 0, "rect": [350, 700, 200, 60]}, [samples["english"], p12])
    ctx2, _ = run("signatures_verify", {}, [outs[0]])
    # the op self-verifies and falls back to an invisible signature when needed
    assert ctx2.result["signatures"][0]["valid"]
