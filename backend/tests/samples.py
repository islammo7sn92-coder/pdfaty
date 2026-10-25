"""Builds the sample documents used by the backend tests and the browser (e2e) tests.

Usage: python3 tests/samples.py <output_dir>   (prints a JSON map name -> path)"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

import pikepdf
from PIL import Image, ImageDraw, ImageFilter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


ARABIC_HTML = """<html><head><meta charset="utf-8"><style>body{{font-family:'Noto Naskh Arabic','Noto Sans Arabic';direction:rtl}}
table{{border-collapse:collapse}} td,th{{border:1px solid #000;padding:4px}}</style></head><body dir="rtl">
<h1>عقد توريد أجهزة</h1>
<p>البند 1.1: يلتزم المورد بتوريد الأجهزة خلال مدة {days} يوما من تاريخ 15/03/2025.</p>
<p>البند 2.3: قيمة العقد الإجمالية {amount} جنيه مصري شاملة الضرائب بنسبة {tax}%.</p>
<p>الرقم القومي للممثل 29801011234567 والبريد sales@example.com</p>
<table><tr><th>الصنف</th><th>الكمية</th><th>السعر</th></tr>
<tr><td>حاسب محمول</td><td>10</td><td>{price}</td></tr><tr><td>طابعة</td><td>5</td><td>8500</td></tr></table>
<p>English line: Contract total is {amount} EGP.</p>
</body></html>"""


def _lo_convert(html: str, out_pdf: str) -> None:
    d = tempfile.mkdtemp()
    src = os.path.join(d, "doc.html")
    with open(src, "w", encoding="utf-8") as fh:
        fh.write(html)
    prof = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{prof}", "--convert-to", "pdf:writer_web_pdf_Export",
                    "--outdir", d, src], capture_output=True, timeout=180)
    shutil.move(os.path.join(d, "doc.pdf"), out_pdf)
    shutil.rmtree(d, ignore_errors=True)
    shutil.rmtree(prof, ignore_errors=True)


def build_samples(d):
    from pathlib import Path

    d = Path(d)
    d.mkdir(parents=True, exist_ok=True)
    s = {}
    # English multi-page with table lines and image
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    img = Image.new("RGB", (600, 400), (30, 120, 200))
    ImageDraw.Draw(img).ellipse([100, 50, 500, 350], fill=(250, 200, 0))
    img_path = str(d / "photo.jpg")
    img.save(img_path, quality=95)
    s["photo"] = img_path
    png = Image.new("RGBA", (300, 120), (0, 0, 0, 0))
    ImageDraw.Draw(png).text((10, 40), "SIGNED", fill=(0, 0, 160, 255))
    png_path = str(d / "sig.png")
    png.save(png_path)
    s["sig"] = png_path

    p = str(d / "english.pdf")
    c = canvas.Canvas(p, pagesize=A4)
    for i in range(5):
        c.setFont("Helvetica", 14)
        c.drawString(72, 760, f"Hello world page {i + 1}")
        c.drawString(72, 740, f"Invoice number INV-{1000 + i} total 12,500.00 USD")
        if i == 0:
            c.drawImage(img_path, 72, 400, 300, 200)
            # simple ruled table
            x0, y0 = 72, 300
            for r in range(4):
                c.line(x0, y0 - r * 20, x0 + 300, y0 - r * 20)
            for col in range(4):
                c.line(x0 + col * 100, y0, x0 + col * 100, y0 - 60)
            for r, row in enumerate([("Item", "Qty", "Price"), ("Pen", "3", "1.50"), ("Book", "2", "12.00")]):
                for col, v in enumerate(row):
                    c.drawString(x0 + col * 100 + 5, y0 - r * 20 - 14, v)
        if i == 3:
            pass  # page 4 has only the header lines
        c.showPage()
    c.showPage()  # page 6 blank
    c.save()
    s["english"] = p

    _lo_convert(ARABIC_HTML.format(days=30, amount="1,250,000", tax=14, price="25000"), str(d / "arabic_v1.pdf"))
    _lo_convert(ARABIC_HTML.format(days=45, amount="1,350,000", tax=14, price="27000"), str(d / "arabic_v2.pdf"))
    s["arabic"] = str(d / "arabic_v1.pdf")
    s["arabic_v2"] = str(d / "arabic_v2.pdf")

    # scanned version of the Arabic page (image only, slightly skewed + noisy)
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(s["arabic"])
    page = doc[0]
    im = page.render(scale=250 / 72).to_pil().convert("L")
    page.close()
    doc.close()
    im = im.rotate(1.2, fillcolor=255, resample=Image.BICUBIC).filter(ImageFilter.GaussianBlur(0.6))
    scan_png = str(d / "scan.png")
    im.save(scan_png)
    s["scan_png"] = scan_png
    from app.ops.convert import images_to_pdf_file

    scan_pdf = str(d / "scanned.pdf")
    images_to_pdf_file([scan_png], scan_pdf, {"size": "a4"})
    s["scanned"] = scan_pdf

    # encrypted
    enc = str(d / "encrypted.pdf")
    with pikepdf.open(p) as pdf:
        pdf.save(enc, encryption=pikepdf.Encryption(user="user123", owner="owner456", R=6))
    s["encrypted"] = enc
    restricted = str(d / "restricted.pdf")
    with pikepdf.open(p) as pdf:
        pdf.save(restricted, encryption=pikepdf.Encryption(user="", owner="owner456", R=6,
                                                           allow=pikepdf.Permissions(modify_other=False, extract=False)))
    s["restricted"] = restricted

    # form
    form = str(d / "form.pdf")
    c = canvas.Canvas(form, pagesize=A4)
    c.drawString(72, 760, "Name:")
    c.acroForm.textfield(name="name", x=130, y=750, width=250, height=20, borderWidth=1)
    c.drawString(72, 720, "Agree:")
    c.acroForm.checkbox(name="agree", x=130, y=715, size=16)
    c.drawString(72, 680, "Gender:")
    c.acroForm.radio(name="gender", value="male", selected=False, x=130, y=675, size=14)
    c.acroForm.radio(name="gender", value="female", selected=False, x=170, y=675, size=14)
    c.acroForm.choice(name="city", options=["Cairo", "Alexandria", "Giza"], value="Cairo", x=130, y=630, width=150, height=20)
    c.showPage()
    c.save()
    s["form"] = form

    bad = str(d / "bad.pdf")
    with open(bad, "wb") as fh:
        fh.write(b"%PDF-1.7\n this is not really a pdf \x00\x01")
    s["bad"] = bad
    docx = str(d / "letter.docx")
    from docx import Document

    dd = Document()
    dd.add_heading("خطاب رسمي", 1)
    dd.add_paragraph("السادة / شركة المثال المحترمين، تحية طيبة وبعد. Hello from DOCX.")
    dd.save(docx)
    s["docx"] = docx
    xlsx = str(d / "sheet.xlsx")
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["البند", "القيمة"])
    wb.active.append(["أجهزة", 1000])
    wb.save(xlsx)
    s["xlsx"] = xlsx
    txt = str(d / "notes.txt")
    with open(txt, "w", encoding="utf-8") as fh:
        fh.write("ملاحظات الاجتماع\nNotes line two\n")
    s["txt"] = txt
    from app.ops.signing import make_test_certificate

    cert = str(d / "test-cert.p12")
    make_test_certificate(cert, "certpw", "PDFATY Test Signer")
    s["cert"] = cert
    return s

if __name__ == "__main__":
    print(json.dumps(build_samples(sys.argv[1]), ensure_ascii=False))
