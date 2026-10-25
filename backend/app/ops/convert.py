"""Conversions to and from PDF."""
from __future__ import annotations

import csv
import io
import json
import os
import shutil
import subprocess
import tempfile
import zipfile

import pikepdf
import pypdfium2 as pdfium
from PIL import Image, ImageOps
from pikepdf import Array, Dictionary, Name

from ..config import soffice_path
from .base import Context, UserError, op, parse_pages
from .extract import geometries, lines_text, page_chars, page_lines, page_tables, build_lines
from .geom import make_image_xobject
from .pages import PAPER, fix_acroform, _merge_acroform_resources, _materialize_inherited


def bundle(ctx: Context, name: str) -> None:
    """Add a ZIP containing all current outputs (when there is more than one)."""
    if len(ctx.outputs) < 2:
        return
    path = ctx.out_path(name)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for o in ctx.outputs:
            z.write(os.path.join(ctx.out_dir, o["file"]), o["name"])
    ctx.add_output(path, name)


# --------------------------------------------------------------------------- images -> PDF
def images_to_pdf_file(paths: list[str], out_path: str, opts: dict) -> int:
    pdf = pikepdf.new()
    size_name = str(opts.get("size", "fit")).lower()
    orientation = opts.get("orientation", "auto")
    margin = float(opts.get("margin", 0))
    enhance = bool(opts.get("enhance"))
    quality = int(opts.get("quality", 85))
    for path in paths:
        with Image.open(path) as im0:
            frames = []
            try:
                n_frames = getattr(im0, "n_frames", 1)
            except Exception:  # noqa: BLE001
                n_frames = 1
            for fi in range(min(n_frames, 500)):
                if n_frames > 1:
                    im0.seek(fi)
                frames.append(im0.copy())
            fmt = im0.format
            dpi = im0.info.get("dpi", (96, 96))
            try:
                orient = im0.getexif().get(274, 1)
            except Exception:  # noqa: BLE001
                orient = 1
        for im in frames:
            im = ImageOps.exif_transpose(im)
            if enhance:
                im = ImageOps.autocontrast(ImageOps.grayscale(im), cutoff=1)
            try:
                dx = float(dpi[0]) if dpi and dpi[0] and float(dpi[0]) > 10 else 96.0
            except (TypeError, ValueError):
                dx = 96.0
            iw, ih = im.size
            if size_name == "fit":
                pw, ph = iw * 72 / dx + 2 * margin, ih * 72 / dx + 2 * margin
            else:
                pw, ph = PAPER.get(size_name, PAPER["a4"])
                if orientation == "landscape" or (orientation == "auto" and iw > ih):
                    pw, ph = max(pw, ph), min(pw, ph)
            aw, ah = pw - 2 * margin, ph - 2 * margin
            s = min(aw / iw, ah / ih)
            dw, dh = iw * s, ih * s
            x, y = (pw - dw) / 2, (ph - dh) / 2
            if fmt == "JPEG" and not enhance and im.mode in ("RGB", "L") and len(frames) == 1 and orient in (1, None):
                with open(path, "rb") as fh:
                    raw = fh.read()
                xo = pikepdf.Stream(pdf, raw)
                xo.Type, xo.Subtype, xo.Filter = Name.XObject, Name.Image, Name.DCTDecode
                xo.Width, xo.Height, xo.BitsPerComponent = iw, ih, 8
                xo.ColorSpace = Name.DeviceRGB if im.mode == "RGB" else Name.DeviceGray
            else:
                if opts.get("jpeg") and im.mode not in ("RGBA", "LA", "P"):
                    buf = io.BytesIO()
                    im.convert("RGB" if im.mode != "L" else "L").save(buf, "JPEG", quality=quality)
                    buf.seek(0)
                    jim = Image.open(buf)
                    xo = make_image_xobject(pdf, jim, quality)
                else:
                    xo = make_image_xobject(pdf, im, quality)
            content = f"q {dw:.3f} 0 0 {dh:.3f} {x:.3f} {y:.3f} cm /Im0 Do Q".encode()
            page = pdf.make_indirect(Dictionary(
                Type=Name.Page, MediaBox=Array([0, 0, pw, ph]),
                Resources=Dictionary(XObject=Dictionary(Im0=xo)), Contents=pdf.make_stream(content)))
            pdf.pages.append(pikepdf.Page(page))
    pdf.docinfo["/Producer"] = "PDFATY"
    pdf.save(out_path, compress_streams=True)
    return len(pdf.pages)


@op("images_to_pdf")
def images_to_pdf(ctx: Context) -> None:
    imgs = [f for f in ctx.inputs if f.kind == "image"]
    if not imgs:
        raise UserError("image_required")
    if ctx.p("separate"):
        for k, f in enumerate(imgs):
            path = ctx.out_path(f.stem + ".pdf")
            images_to_pdf_file([f.path], path, ctx.params)
            ctx.add_output(path)
            ctx.progress(100 * (k + 1) / len(imgs))
        bundle(ctx, "images_pdf.zip")
        return
    name = (ctx.p("output_name") or imgs[0].stem) + ".pdf"
    path = ctx.out_path(name)
    n = images_to_pdf_file([f.path for f in imgs], path, ctx.params)
    ctx.add_output(path, name)
    ctx.result = {"pages": n}


# --------------------------------------------------------------------------- office / text / html -> PDF
_EXPORT_FILTER = {
    "writer": "writer_pdf_Export", "calc": "calc_pdf_Export", "impress": "impress_pdf_Export", "draw": "draw_pdf_Export",
}


def _lo_app(ext: str) -> str:
    if ext in (".xls", ".xlsx", ".ods", ".csv", ".tsv"):
        return "calc"
    if ext in (".ppt", ".pptx", ".odp"):
        return "impress"
    if ext in (".odg", ".vsd", ".vsdx", ".svg", ".dxf", ".cdr", ".eps", ".pub"):
        return "draw"
    return "writer"


def libreoffice_to_pdf(src: str, out_dir: str, pdfa: bool = False, timeout: int = 240) -> str:
    exe = soffice_path()
    if not exe:
        raise UserError("office_unavailable")
    ext = os.path.splitext(src)[1].lower()
    app = _lo_app(ext)
    flt = _EXPORT_FILTER[app]
    opts = {"ExportFormFields": {"type": "boolean", "value": "true"}}
    if pdfa:
        opts["SelectPdfVersion"] = {"type": "long", "value": "2"}  # PDF/A-2b
    target = f"pdf:{flt}:{json.dumps(opts)}"
    profile = tempfile.mkdtemp(prefix="lo-profile-")
    cmd = [exe, "--headless", "--norestore", "--nolockcheck", "--nodefault", "--nologo",
           f"-env:UserInstallation=file://{profile}", "--convert-to", target, "--outdir", out_dir]
    if ext in (".txt", ".md", ".log", ".json", ".xml"):
        cmd += ["--infilter=Text (encoded):UTF8,LF,,,"]
    elif ext == ".csv":
        cmd += ["--infilter=CSV:44,34,76,1"]
    elif ext == ".tsv":
        cmd += ["--infilter=CSV:9,34,76,1"]
    cmd.append(src)
    env = dict(os.environ, HOME=profile)
    try:
        subprocess.run(cmd, capture_output=True, timeout=timeout, env=env, check=False)
    except subprocess.TimeoutExpired:
        raise UserError("timeout")
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = os.path.join(out_dir, os.path.splitext(os.path.basename(src))[0] + ".pdf")
    if not os.path.exists(out) or os.path.getsize(out) == 0:
        raise UserError("conversion_failed")
    return out


@op("convert_to_pdf")
def convert_to_pdf(ctx: Context) -> None:
    if not ctx.inputs:
        raise UserError("no_files")
    produced: list[tuple[str, str]] = []
    pdfa = bool(ctx.p("pdfa"))
    for k, f in enumerate(ctx.inputs):
        work = tempfile.mkdtemp(dir=os.path.dirname(ctx.tmp("x")))
        if f.kind == "pdf":
            produced.append((f.path, f.stem))
        elif f.kind == "image":
            path = os.path.join(work, "img.pdf")
            images_to_pdf_file([f.path], path, ctx.params)
            produced.append((path, f.stem))
        elif f.kind in ("office", "text", "html"):
            # give LibreOffice a clean name with the right extension
            ext = os.path.splitext(f.name)[1].lower() or ".txt"
            src = os.path.join(work, "document" + ext)
            shutil.copy(f.path, src)
            produced.append((libreoffice_to_pdf(src, work, pdfa), f.stem))
        else:
            raise UserError("unsupported_file_type", f.name)
        ctx.progress(90 * (k + 1) / len(ctx.inputs))
    names = ctx.p("names") or []
    if ctx.p("combine", len(produced) > 1) and len(produced) > 1:
        out = pikepdf.new()
        for path, stem in produced:
            src = pikepdf.open(path)
            _merge_acroform_resources(out, src)
            start = len(out.pages)
            for p in src.pages:
                _materialize_inherited(p)
                out.pages.append(p)
            with out.open_outline() as ol:
                ol.root.append(pikepdf.OutlineItem(stem, start))
        fix_acroform(out)
        ctx.save_pdf(out, (ctx.p("output_name") or "combined") + ".pdf")
    else:
        for k, (path, stem) in enumerate(produced):
            name = (names[k] if k < len(names) and names[k] else stem)
            name = name if name.lower().endswith(".pdf") else name + ".pdf"
            dest = ctx.out_path(name)
            shutil.copy(path, dest)
            ctx.add_output(dest, name)
        bundle(ctx, "converted.zip")


# --------------------------------------------------------------------------- PDF -> images / text
def _open_pdfium(ctx: Context):
    f = ctx.first_pdf()
    ctx.open_pdf(f).close()  # validates password + owner permissions
    pw = ctx.password_for(ctx.inputs.index(f))
    ctx.result["_geos"] = geometries(f.path, pw)
    try:
        return f, pdfium.PdfDocument(f.path, password=pw or None)
    except pdfium.PdfiumError as e:
        if "password" in str(e).lower():
            raise UserError("wrong_password" if ctx.password_for(0) else "password_required")
        raise UserError("invalid_pdf")


@op("pdf_to_images")
def pdf_to_images(ctx: Context) -> None:
    f, doc = _open_pdfium(ctx)
    fmt = "jpg" if str(ctx.p("format", "png")).lower() in ("jpg", "jpeg") else "png"
    dpi = ctx.pint("dpi", 150, 36, 600)
    quality = ctx.pint("quality", 88, 30, 100)
    pages = parse_pages(ctx.p("pages"), len(doc))
    for k, i in enumerate(pages):
        page = doc[i]
        scale = dpi / 72
        w, h = page.get_size()
        if w * h * scale * scale > 60_000_000:
            scale = (60_000_000 / (w * h)) ** 0.5
        img = page.render(scale=scale, may_draw_forms=True).to_pil()
        page.close()
        name = f"{f.stem}_p{i + 1}.{fmt}"
        path = ctx.out_path(name)
        if fmt == "jpg":
            img.convert("RGB").save(path, "JPEG", quality=quality, dpi=(dpi, dpi))
        else:
            img.save(path, "PNG", dpi=(dpi, dpi), optimize=False)
        ctx.add_output(path, name)
        ctx.progress(100 * (k + 1) / len(pages))
    doc.close()
    bundle(ctx, f"{f.stem}_images.zip")


def pdf_text_pages(doc: pdfium.PdfDocument, pages: list[int], geos=None) -> list[str]:
    out = []
    for i in pages:
        page = doc[i]
        out.append(lines_text(page_lines(page, geo=geos[i] if geos else None)))
        page.close()
    return out


@op("pdf_to_text")
def pdf_to_text(ctx: Context) -> None:
    f, doc = _open_pdfium(ctx)
    pages = parse_pages(ctx.p("pages"), len(doc))
    texts = pdf_text_pages(doc, pages, ctx.result.pop("_geos", None))
    doc.close()
    body = "\n\n".join(f"--- {p + 1} ---\n{t}" if ctx.p("page_markers", True) else t for p, t in zip(pages, texts))
    path = ctx.out_path(f.stem + ".txt")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(body)
    ctx.add_output(path)
    empty = [p + 1 for p, t in zip(pages, texts) if not t.strip()]
    ctx.result = {"chars": len(body), "pages_without_text": empty}


# --------------------------------------------------------------------------- PDF -> DOCX
def _set_rtl(paragraph, rtl: bool) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if not rtl:
        return
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    ppr.insert(0, bidi)
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rtl_el = OxmlElement("w:rtl")
        rpr.append(rtl_el)


def _style_run(run, size: float, bold: bool, font: str = "Arial") -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.size = Pt(max(6, min(72, round(size))))
    run.font.bold = bold
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        rfonts.set(qn("w:cs"), font)
    from docx.oxml import OxmlElement

    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(max(6, min(72, round(size))) * 2)))
    rpr.append(szcs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))


@op("pdf_to_docx")
def pdf_to_docx(ctx: Context) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    f, doc = _open_pdfium(ctx)
    geos = ctx.result.pop("_geos")
    pw = ctx.password_for(ctx.inputs.index(f))
    pages = parse_pages(ctx.p("pages"), len(doc))
    out = Document()
    sec = out.sections[0]
    w0, h0 = geos[pages[0]].vw, geos[pages[0]].vh
    sec.page_width, sec.page_height = Pt(w0), Pt(h0)
    if w0 > h0:
        sec.orientation = WD_ORIENT.LANDSCAPE
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Pt(42))
    normal = out.styles["Normal"]
    normal.font.name = "Arial"
    scanned = []
    for k, i in enumerate(pages):
        page = doc[i]
        tp = page.get_textpage()
        height = page.get_height()
        geo = geos[i]
        chars = page_chars(page, tp, geo)
        lines = build_lines(chars, page_width=geo.vw)
        tables = page_tables(f.path, pw, i, chars) if lines and geo.rot == 0 else []
        if not lines:
            scanned.append(i + 1)
            img = page.render(scale=150 / 72).to_pil().convert("RGB")
            buf = io.BytesIO()
            img.save(buf, "JPEG", quality=85)
            buf.seek(0)
            out.add_picture(buf, width=sec.page_width - sec.left_margin - sec.right_margin)
        else:
            blocks: list[tuple[float, str, object]] = []
            for t in tables:
                blocks.append((t["bbox"][1], "table", t))

            def in_table(ln):
                cy = (ln.t + ln.b) / 2
                cx = (ln.l + ln.r) / 2
                return any(t["bbox"][0] - 2 <= cx <= t["bbox"][2] + 2 and t["bbox"][1] - 2 <= cy <= t["bbox"][3] + 2 for t in tables)

            para: list = []
            prev = None
            paragraphs = []
            for ln in sorted([l for l in lines if not in_table(l)], key=lambda l: (round(l.t, 0), -l.r if l.rtl else l.l)):
                if prev is not None and (ln.t - prev.b > max(prev.size, ln.size) * 0.9 or abs(ln.size - prev.size) > 1.5
                                         or ln.rtl != prev.rtl or ln.bold != prev.bold):
                    paragraphs.append(para)
                    para = []
                para.append(ln)
                prev = ln
            if para:
                paragraphs.append(para)
            for p in paragraphs:
                blocks.append((p[0].t, "para", p))
            blocks.sort(key=lambda b: b[0])
            content_w = w0 - 84
            for _, kind, obj in blocks:
                if kind == "table":
                    rows = obj["rows"]  # type: ignore[index]
                    ncols = max(len(r) for r in rows)
                    table = out.add_table(rows=len(rows), cols=ncols)
                    table.style = "Table Grid"
                    rtl_table = any(any(c for c in r if c and any("؀" <= ch <= "ۿ" for ch in c)) for r in rows)
                    for ri, r in enumerate(rows):
                        cells = list(r) + [""] * (ncols - len(r))
                        for ci, val in enumerate(cells):
                            cell = table.cell(ri, ci)
                            cp = cell.paragraphs[0]
                            run = cp.add_run(val)
                            _style_run(run, 10, ri == 0)
                            if rtl_table:
                                cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                                _set_rtl(cp, True)
                    if rtl_table:
                        from docx.oxml import OxmlElement

                        tblpr = table._tbl.tblPr
                        tblpr.append(OxmlElement("w:bidiVisual"))
                    out.add_paragraph()
                else:
                    p = obj  # type: ignore[assignment]
                    rtl = p[0].rtl
                    text = " ".join(l.text for l in p)
                    par = out.add_paragraph()
                    run = par.add_run(text)
                    _style_run(run, p[0].size, p[0].bold)
                    left = min(l.l for l in p)
                    right = max(l.r for l in p)
                    width = right - left
                    centred = abs((left - 42) - (w0 - 42 - right)) < 20 and width < content_w * 0.75
                    if centred:
                        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif rtl:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    par.paragraph_format.space_after = Pt(4)
                    _set_rtl(par, rtl)
        tp.close()
        page.close()
        if k < len(pages) - 1:
            out.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        ctx.progress(100 * (k + 1) / len(pages))
    doc.close()
    path = ctx.out_path(f.stem + ".docx")
    out.save(path)
    ctx.add_output(path)
    ctx.result = {"scanned_pages": scanned}


# --------------------------------------------------------------------------- PDF -> XLSX / CSV
def _collect_tables(ctx: Context, f, doc, pages) -> list[tuple[int, list[list[str]]]]:
    pw = ctx.password_for(ctx.inputs.index(f))
    geos = ctx.result.pop("_geos", None) or geometries(f.path, pw)
    found = []
    for k, i in enumerate(pages):
        page = doc[i]
        tp = page.get_textpage()
        geo = geos[i]
        chars = page_chars(page, tp, geo)
        if geo.rot == 0:
            for t in page_tables(f.path, pw, i, chars):
                found.append((i, t["rows"]))
        if ctx.p("fallback_lines", True) and not any(p == i for p, _ in found):
            lines = build_lines(chars, page_width=geo.vw)
            if lines:
                found.append((i, [[l.text] for l in lines]))
        tp.close()
        page.close()
        ctx.progress(90 * (k + 1) / len(pages))
    return found


def _cell_value(v: str):
    s = v.strip()
    norm = s.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩٫٬", "0123456789.,")).replace(",", "")
    try:
        if norm and norm.replace(".", "", 1).replace("-", "", 1).isdigit():
            return float(norm) if "." in norm else int(norm)
    except ValueError:
        pass
    return s


@op("pdf_to_xlsx")
def pdf_to_xlsx(ctx: Context) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    f, doc = _open_pdfium(ctx)
    pages = parse_pages(ctx.p("pages"), len(doc))
    tables = _collect_tables(ctx, f, doc, pages)
    doc.close()
    wb = Workbook()
    wb.remove(wb.active)
    if not tables:
        ws = wb.create_sheet("Empty")
        ws["A1"] = "No text found"
    per_page = bool(ctx.p("sheet_per_table", True))
    ws = None
    counter: dict[int, int] = {}
    for page_i, rows in tables:
        if per_page or ws is None:
            counter[page_i] = counter.get(page_i, 0) + 1
            ws = wb.create_sheet(f"P{page_i + 1}-{counter[page_i]}"[:31])
            rtl = any(any("؀" <= ch <= "ۿ" for ch in c) for r in rows for c in r)
            ws.sheet_view.rightToLeft = rtl
        start = ws.max_row + (2 if ws.max_row > 1 else 0)
        for ri, r in enumerate(rows):
            for ci, v in enumerate(r):
                cell = ws.cell(row=start + ri, column=ci + 1, value=_cell_value(v))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if ri == 0 and len(r) > 1:
                    cell.font = Font(bold=True)
        for col in ws.columns:
            width = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = max(8, min(60, width + 2))
    path = ctx.out_path(f.stem + ".xlsx")
    wb.save(path)
    ctx.add_output(path)
    ctx.result = {"tables": sum(1 for _, r in tables if r and len(r[0]) > 1)}


@op("pdf_to_csv")
def pdf_to_csv(ctx: Context) -> None:
    f, doc = _open_pdfium(ctx)
    pages = parse_pages(ctx.p("pages"), len(doc))
    ctx.params.setdefault("fallback_lines", False)
    tables = _collect_tables(ctx, f, doc, pages)
    doc.close()
    if not tables:
        raise UserError("no_tables_found")
    for n, (page_i, rows) in enumerate(tables):
        name = f"{f.stem}_p{page_i + 1}_table{n + 1}.csv"
        path = ctx.out_path(name)
        with open(path, "w", encoding="utf-8-sig", newline="") as fh:
            csv.writer(fh).writerows(rows)
        ctx.add_output(path, name)
    bundle(ctx, f"{f.stem}_tables.zip")
    ctx.result = {"tables": len(tables)}


# --------------------------------------------------------------------------- PDF -> PPTX
@op("pdf_to_pptx")
def pdf_to_pptx(ctx: Context) -> None:
    from pptx import Presentation
    from pptx.util import Emu, Pt

    f, doc = _open_pdfium(ctx)
    geos = ctx.result.pop("_geos")
    pages = parse_pages(ctx.p("pages"), len(doc))
    mode = ctx.p("mode", "image")
    dpi = ctx.pint("dpi", 150, 72, 300)
    prs = Presentation()
    w0, h0 = geos[pages[0]].vw, geos[pages[0]].vh
    EMU = 12700
    prs.slide_width, prs.slide_height = Emu(int(w0 * EMU)), Emu(int(h0 * EMU))
    blank = prs.slide_layouts[6]
    for k, i in enumerate(pages):
        page = doc[i]
        slide = prs.slides.add_slide(blank)
        lines = page_lines(page, geo=geos[i])
        sx = prs.slide_width / (geos[i].vw * EMU)
        sy = prs.slide_height / (geos[i].vh * EMU)
        if mode == "editable" and lines:
            for ln in lines:
                tb = slide.shapes.add_textbox(Emu(int(ln.l * EMU * sx)), Emu(int(ln.t * EMU * sy)),
                                              Emu(int(max(10, ln.r - ln.l + 4) * EMU * sx)), Emu(int(max(8, ln.b - ln.t) * EMU * sy)))
                tf = tb.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = ln.text
                r.font.size = Pt(max(6, ln.size * min(sx, sy)))
                r.font.bold = ln.bold
                if ln.rtl:
                    from pptx.enum.text import PP_ALIGN
                    p.alignment = PP_ALIGN.RIGHT
                    pPr = p._p.get_or_add_pPr()
                    pPr.set("rtl", "1")
        else:
            img = page.render(scale=dpi / 72, may_draw_forms=True).to_pil().convert("RGB")
            buf = io.BytesIO()
            img.save(buf, "JPEG", quality=88)
            buf.seek(0)
            slide.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
        if lines:
            slide.notes_slide.notes_text_frame.text = "\n".join(l.text for l in lines)[:20000]
        page.close()
        ctx.progress(100 * (k + 1) / len(pages))
    doc.close()
    path = ctx.out_path(f.stem + ".pptx")
    prs.save(path)
    ctx.add_output(path)


# --------------------------------------------------------------------------- extract embedded images
@op("extract_images")
def extract_images(ctx: Context) -> None:
    from pikepdf import PdfImage

    f = ctx.first_pdf()
    pdf = ctx.open_pdf(f)
    pages = parse_pages(ctx.p("pages"), len(pdf.pages))
    min_px = ctx.pint("min_size", 32, 1, 10000)
    seen = set()
    count = 0
    for k, i in enumerate(pages):
        page = pdf.pages[i]
        stack = [page.obj.get("/Resources")]
        depth = 0
        while stack and depth < 200:
            depth += 1
            res = stack.pop()
            if res is None or "/XObject" not in res:
                continue
            for _, xo in res.XObject.items():
                if xo.objgen in seen:
                    continue
                seen.add(xo.objgen)
                sub = xo.get("/Subtype")
                if sub == Name.Form:
                    stack.append(xo.get("/Resources"))
                    continue
                if sub != Name.Image or int(xo.get("/Width", 0)) < min_px or int(xo.get("/Height", 0)) < min_px:
                    continue
                if xo.get("/ImageMask"):
                    continue
                try:
                    pim = PdfImage(xo)
                    base = os.path.join(ctx.out_dir, f"{f.stem}_p{i + 1}_img{count + 1}")
                    outp = pim.extract_to(fileprefix=base)
                except Exception:  # noqa: BLE001 – unsupported encodings: render through PIL
                    try:
                        im = PdfImage(xo).as_pil_image()
                        outp = os.path.join(ctx.out_dir, f"{f.stem}_p{i + 1}_img{count + 1}.png")
                        im.save(outp)
                    except Exception:  # noqa: BLE001
                        continue
                count += 1
                ctx.add_output(outp)
                if count >= 2000:
                    break
        ctx.progress(100 * (k + 1) / len(pages))
    if not count:
        raise UserError("no_images_found")
    bundle(ctx, f"{f.stem}_images.zip")
    ctx.result = {"images": count}
